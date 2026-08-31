#!/usr/bin/env python3
"""Export the four edited manuscript sections to a review-ready DOCX."""

from __future__ import annotations

import argparse
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPO_ROOT = Path(__file__).resolve().parents[1]
SOURCE_TEX = REPO_ROOT / "manuscript.tex"
BIBLIOGRAPHY = REPO_ROOT / "bibliography" / "references.bib"
CSL = REPO_ROOT / "tools" / "acs-numeric.csl"

SECTION_HEADINGS = (
    "Introduction",
    "Materials and Methods",
    "Results",
    "Discussion",
    "References",
)

FIGURE_REFERENCES = {
    r"Figure~\ref{fig:descriptive_bi}": "Figure 2",
    r"Figure~\ref{fig:thirty_minute_response}": "Figure 3",
    r"Figure~\ref{fig:interaction_wind_sun_nextday}": "Figure 4",
}


def extract_body(source: str) -> str:
    start_marker = r"\section{Introduction}"
    end_marker = r"\section{Conclusions}"
    try:
        start = source.index(start_marker)
        end = source.index(end_marker, start)
    except ValueError as exc:
        raise RuntimeError("Could not locate the requested manuscript section boundaries") from exc

    body = source[start:end].rstrip()
    body = body.replace(
        r"Appendix~\ref{app:deployment_metadata}",
        "Appendix A",
    )
    body = body.replace(
        r"Appendices~\ref{app:30min_models} and \ref{app:nextday_models}",
        "Appendices B and C",
    )
    for latex, word_text in FIGURE_REFERENCES.items():
        body = body.replace(latex, word_text)
    body = body.replace(r"$^{\circ}$", "°")

    if r"\ref{" in body:
        raise RuntimeError("An unresolved LaTeX cross-reference remains in the export body")
    return body


def write_latex_wrapper(path: Path, body: str) -> None:
    path.write_text(
        "\n".join(
            (
                r"\documentclass{article}",
                r"\usepackage{amsmath}",
                r"\usepackage{graphicx}",
                r"\usepackage{float}",
                r"\usepackage{hyperref}",
                r"\begin{document}",
                body,
                r"\end{document}",
                "",
            )
        ),
        encoding="utf-8",
    )


def set_run_font(run, name: str, size: float, color: str = "000000", bold=None, italic=None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), name)


def set_style_font(style, name: str, size: float, color: str, bold=None, italic=None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), name)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    set_run_font(run, "Calibri", 9, "666666")
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend((fld_char_begin, instr_text, fld_char_end))


def ensure_update_fields(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def style_document(doc: Document) -> None:
    # Narrative proposal preset. Named overrides are caption and bibliography
    # styling, page numbers, and page breaks before major sections after the
    # Introduction.
    section = doc.sections[0]
    section.start_type = WD_SECTION.CONTINUOUS
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    for style_name in ("Normal", "Body Text", "First Paragraph"):
        style = doc.styles[style_name]
        set_style_font(style, "Calibri", 11, "000000")
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.333
        style.paragraph_format.widow_control = True

    heading_specs = {
        "Heading 1": (16, "2E74B5", 18, 10),
        "Heading 2": (13, "2E74B5", 12, 6),
        "Heading 3": (12, "1F4D78", 8, 4),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = doc.styles[style_name]
        set_style_font(style, "Calibri", size, color, bold=True)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption_style = doc.styles["Image Caption"]
    set_style_font(caption_style, "Calibri", 9, "333333", italic=False)
    caption_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    caption_style.paragraph_format.space_before = Pt(6)
    caption_style.paragraph_format.space_after = Pt(12)
    caption_style.paragraph_format.line_spacing = 1.0
    caption_style.paragraph_format.keep_together = True

    bibliography_style = doc.styles["Bibliography"]
    set_style_font(bibliography_style, "Calibri", 10, "000000")
    bibliography_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    bibliography_style.paragraph_format.left_indent = Inches(0.25)
    bibliography_style.paragraph_format.first_line_indent = Inches(-0.25)
    bibliography_style.paragraph_format.space_before = Pt(0)
    bibliography_style.paragraph_format.space_after = Pt(6)
    bibliography_style.paragraph_format.line_spacing = 1.15
    bibliography_style.paragraph_format.widow_control = True

    expected_figure_number = 1
    first_heading = True
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        style_name = paragraph.style.name if paragraph.style else ""

        if style_name == "Heading 1":
            paragraph.paragraph_format.page_break_before = not first_heading
            first_heading = False

        if style_name == "Image Caption" and text:
            if not text.lower().startswith("figure "):
                paragraph.text = f"Figure {expected_figure_number}. {text}"
            expected_figure_number += 1
            for run in paragraph.runs:
                set_run_font(run, "Calibri", 9, "333333")

        if "m:oMathPara" in paragraph._p.xml:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.keep_together = True

        if "w:drawing" in paragraph._p.xml:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True

    if expected_figure_number != 5:
        raise RuntimeError(f"Expected four figure captions, found {expected_figure_number - 1}")

    for shape in doc.inline_shapes:
        max_width = Inches(6.35)
        max_height = Inches(7.25)
        scale = min(1.0, max_width / shape.width, max_height / shape.height)
        if scale < 1.0:
            shape.width = int(shape.width * scale)
            shape.height = int(shape.height * scale)

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.clear()
    add_page_number(footer_paragraph)
    ensure_update_fields(doc)

    doc.core_properties.title = "Monarch wind and light manuscript substantive sections"
    doc.core_properties.subject = "Co-author review copy"
    doc.core_properties.author = "Kyle Nessen"
    doc.core_properties.keywords = "monarch butterfly, wind, light, manuscript"


def validate_document(path: Path) -> None:
    doc = Document(path)
    headings = [p.text.strip() for p in doc.paragraphs if p.style and p.style.name == "Heading 1"]
    if tuple(headings) != SECTION_HEADINGS:
        raise RuntimeError(f"Unexpected top-level headings: {headings}")

    full_text = "\n".join(p.text for p in doc.paragraphs)
    excluded = ("Simple Summary", "Abstract", "Conclusions")
    present = [label for label in excluded if label in full_text]
    if present:
        raise RuntimeError(f"Excluded section labels remain: {present}")
    if r"\ref{" in full_text or r"\cite{" in full_text:
        raise RuntimeError("Raw LaTeX references remain in the Word document")
    if len(doc.inline_shapes) != 4:
        raise RuntimeError(f"Expected four figures, found {len(doc.inline_shapes)}")
    bibliography_entries = [
        p for p in doc.paragraphs if p.style and p.style.name == "Bibliography"
    ]
    if len(bibliography_entries) != 37:
        raise RuntimeError(
            f"Expected 37 cited references, found {len(bibliography_entries)}"
        )
    if not any("[1" in p.text for p in doc.paragraphs):
        raise RuntimeError("Numbered in-text citations were not generated")


def run_export(output: Path, pandoc: str) -> None:
    source = SOURCE_TEX.read_text(encoding="utf-8")
    body = extract_body(source)
    output.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory(prefix="monarch-docx-") as tmp_name:
        tmp_dir = Path(tmp_name)
        wrapped_tex = tmp_dir / "substantive-sections.tex"
        raw_docx = tmp_dir / "substantive-sections-raw.docx"
        write_latex_wrapper(wrapped_tex, body)

        subprocess.run(
            (
                pandoc,
                str(wrapped_tex),
                "--from=latex",
                "--to=docx",
                "--citeproc",
                f"--csl={CSL}",
                f"--bibliography={BIBLIOGRAPHY}",
                f"--resource-path={REPO_ROOT}",
                "--metadata=link-citations:false",
                "--metadata=reference-section-title:References",
                f"--output={raw_docx}",
            ),
            cwd=REPO_ROOT,
            check=True,
        )

        doc = Document(raw_docx)
        style_document(doc)
        doc.save(output)

    validate_document(output)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--pandoc", default=shutil.which("pandoc") or "pandoc")
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    run_export(args.output.resolve(), args.pandoc)
