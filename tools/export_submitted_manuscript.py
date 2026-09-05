#!/usr/bin/env python3
"""Export a clean MDPI Insects DOCX from a specific Git commit."""

from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path

from lxml import etree
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


REPO_ROOT = Path(__file__).resolve().parents[1]
REFERENCE_TEMPLATE = REPO_ROOT / "insects-template.dot"
CSL = REPO_ROOT / "tools" / "acs-numeric.csl"
EXPECTED_TEMPLATE_SHA256 = "993079177f01de14a7a9c3be76bf2e1f84abb1e9b53a8cc78b00541943123db0"
EXPECTED_COMMIT = "12b7cc9465ce42000c7143ad6d3e10c7311aee5c"
TEMPLATE_MAIN_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
)
DOCUMENT_MAIN_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
)


@dataclass
class TableSpec:
    marker: str
    caption: str
    label: str | None
    rows: list[list[str]]
    merge_rows: set[int]


def run(*args: str, cwd: Path | None = None, capture: bool = False) -> str:
    result = subprocess.run(
        args,
        cwd=cwd,
        check=True,
        text=True,
        stdout=subprocess.PIPE if capture else None,
    )
    return result.stdout.strip() if capture else ""


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def extract_command(text: str, command: str) -> str:
    start = text.index(f"\\{command}")
    brace = text.index("{", start)
    depth = 0
    for index in range(brace, len(text)):
        char = text[index]
        if char == "{":
            depth += 1
        elif char == "}":
            depth -= 1
            if depth == 0:
                return text[brace + 1 : index]
    raise RuntimeError(f"Unclosed command: {command}")


def strip_latex(value: str) -> str:
    value = value.strip()
    value = re.sub(r"(?<!\\)%.*", "", value)
    previous = None
    while value != previous:
        previous = value
        value = re.sub(r"\\(?:textbf|textit|emph|mathrm|textrm|texttt)\{([^{}]*)\}", r"\1", value)
    replacements = {
        r"\&": "&",
        r"\%": "%",
        r"\_": "_",
        r"\quad": " ",
        r"\;": " ",
        r"\,": " ",
        "\\ ": " ",
        r"~": " ",
        r"\Delta": "Δ",
        r"\beta": "β",
        r"\rho": "ρ",
        r"\pm": "±",
        r"\times": "×",
        r"\geq": "≥",
        r"\leq": "≤",
        r"\sim": "~",
        r"\circ": "°",
        r"\min": "min",
        r"\max": "max",
        r"$-$": "−",
        r"---": "-",
        r"--": "-",
    }
    for old, new in replacements.items():
        value = value.replace(old, new)
    value = re.sub(r"\\sqrt\[3\]\{\\Delta(?:\\mathrm\{BI\}|BI)\}", "∛ΔBI", value)
    value = re.sub(r"\\sqrt\{\|\\Delta(?:\\mathrm\{BI\}|\\?BI)_?\{?\\max\}?\|\}", "√|ΔBImax|", value)
    value = re.sub(r"\^\{2\}", "²", value)
    value = value.replace("^2", "²")
    value = re.sub(r"_\{\\text\{([^{}]+)\}\}", r"_\1", value)
    value = re.sub(r"_\{([^{}]+)\}", r"_\1", value)
    value = value.replace("$", "")
    value = re.sub(r"\\(?:small|footnotesize|centering|raggedright|arraybackslash)\b", "", value)
    value = re.sub(r"\\[A-Za-z]+\*?(?:\[[^\]]*\])?", "", value)
    value = value.replace("{", "").replace("}", "")
    value = value.replace("^2", "²")
    return re.sub(r"\s+", " ", value).strip()


def split_table_rows(body: str) -> list[str]:
    rows: list[str] = []
    current: list[str] = []
    index = 0
    math = False
    while index < len(body):
        char = body[index]
        if char == "$" and (index == 0 or body[index - 1] != "\\"):
            math = not math
        if not math and body[index : index + 2] == r"\\":
            rows.append("".join(current))
            current = []
            index += 2
            continue
        current.append(char)
        index += 1
    if "".join(current).strip():
        rows.append("".join(current))
    return rows


def split_cells(row: str) -> list[str]:
    cells: list[str] = []
    current: list[str] = []
    depth = 0
    math = False
    for index, char in enumerate(row):
        if char == "$" and (index == 0 or row[index - 1] != "\\"):
            math = not math
        elif not math and char == "{":
            depth += 1
        elif not math and char == "}":
            depth = max(0, depth - 1)
        if char == "&" and not math and depth == 0 and (index == 0 or row[index - 1] != "\\"):
            cells.append("".join(current))
            current = []
        else:
            current.append(char)
    cells.append("".join(current))
    return cells


def parse_table_block(block: str, number: int) -> TableSpec:
    marker = f"[[TABLE_{number:02d}]]"
    if r"\caption" not in block:
        raise RuntimeError(f"Table {number} has no caption")
    caption = strip_latex(extract_command(block, "caption"))
    label_match = re.search(r"\\label\{([^}]+)\}", block)
    label = label_match.group(1) if label_match else None

    tabular_match = re.search(
        r"\\begin\{(?:tabularx|tabular|longtable)\}(?:\{[^\n]*\}){1,2}(.*?)\\end\{(?:tabularx|tabular|longtable)\}",
        block,
        re.S,
    )
    if not tabular_match:
        raise RuntimeError(f"Could not parse table body for {marker}")
    body = tabular_match.group(1)
    body = re.sub(r"\\caption\{.*?\}(?:\s*\\label\{[^}]+\})?\s*\\\\", "", body, flags=re.S)
    body = re.sub(r"\\endfirsthead.*?\\endfoot", "", body, flags=re.S)
    body = re.sub(r"\\(?:toprule|midrule|bottomrule|endfirsthead|endhead|endfoot|endlastfoot|addlinespace)\b", "", body)
    body = re.sub(r"(?<!\\)%.*", "", body)

    parsed: list[list[str]] = []
    merge_rows: set[int] = set()
    for raw_row in split_table_rows(body):
        raw_row = raw_row.strip()
        if not raw_row:
            continue
        multicol = re.fullmatch(r"\\multicolumn\{\d+\}\{[^}]+\}\{(.*)\}", raw_row, re.S)
        if multicol:
            parsed.append([strip_latex(multicol.group(1))])
            merge_rows.add(len(parsed) - 1)
            continue
        cells = split_cells(raw_row)
        cleaned: list[str] = []
        merge = False
        for cell in cells:
            match = re.fullmatch(r"\\multicolumn\{\d+\}\{[^}]+\}\{(.*)\}", cell.strip(), re.S)
            if match:
                cleaned.append(strip_latex(match.group(1)))
                merge = True
            else:
                cleaned.append(strip_latex(cell))
        if any(cleaned):
            parsed.append(cleaned)
            if merge:
                merge_rows.add(len(parsed) - 1)

    if not parsed:
        raise RuntimeError(f"No rows parsed for {marker}")
    if len(parsed) > 1 and parsed[0] == parsed[1]:
        parsed.pop(1)
        merge_rows = {index - 1 if index > 1 else index for index in merge_rows if index != 1}
    columns = max(len(row) for row in parsed if row)
    rows = [row + [""] * (columns - len(row)) for row in parsed]
    return TableSpec(marker, caption, label, rows, merge_rows)


def extract_tables(source: str) -> tuple[str, list[TableSpec]]:
    tables: list[TableSpec] = []
    pattern = re.compile(r"\\begin\{(table|longtable)\}(?:\[[^\]]*\])?")
    output: list[str] = []
    position = 0
    while match := pattern.search(source, position):
        env = match.group(1)
        end_token = f"\\end{{{env}}}"
        end = source.index(end_token, match.end()) + len(end_token)
        block = source[match.start() : end]
        spec = parse_table_block(block, len(tables) + 1)
        tables.append(spec)
        output.append(source[position : match.start()])
        output.append(f"\n\n{spec.marker}\n\n")
        position = end
    output.append(source[position:])
    return "".join(output), tables


def replace_cross_references(source: str, original: str, tables: list[TableSpec]) -> str:
    figure_labels = re.findall(r"\\label\{(fig:[^}]+)\}", original)
    table_labels = [table.label for table in tables if table.label]
    appendix_maps: dict[str, str] = {}
    appendix_start = original.find(r"\begin{appendix}")
    if appendix_start >= 0:
        appendix_number = 0
        appendix_subsection = 0
        appendix_suffix = ""
        appendix_source = original[appendix_start:]
        for match in re.finditer(
            r"\\(?:sub)?section\{[^{}]+\}|\\label\{(app:[^}]+)\}", appendix_source
        ):
            if match.group(0).startswith(r"\section"):
                appendix_number += 1
                appendix_subsection = 0
                appendix_suffix = ""
            elif match.group(0).startswith(r"\subsection"):
                appendix_subsection += 1
                appendix_suffix = f".{appendix_subsection}"
            elif match.group(1):
                appendix_maps[match.group(1)] = f"Appendix {chr(64 + appendix_number)}{appendix_suffix}"
    revised_tables = "tab:deployment_metadata" in table_labels
    maps = {
        **{label: f"Figure {index}" for index, label in enumerate(figure_labels, 1)},
        **{
            label: f"Table A{index}" if revised_tables else f"Table {index}"
            for index, label in enumerate(table_labels, 1)
        },
        **appendix_maps,
    }

    source = re.sub(
        r"(?:Figure|Table|Appendix)~\\ref\{([^}]+)\}",
        lambda match: maps[match.group(1)],
        source,
    )
    source = re.sub(
        r"Appendices~\\ref\{([^}]+)\}\s+and\s+\\ref\{([^}]+)\}",
        lambda match: (
            maps[match.group(1)]
            if maps[match.group(1)] == maps[match.group(2)]
            else f"{maps[match.group(1)]} and {maps[match.group(2)]}"
        ),
        source,
    )
    source = re.sub(r"\\label\{(?:tab|app):[^}]+\}", "", source)
    unresolved = re.findall(r"\\ref\{([^}]+)\}", source)
    if unresolved:
        raise RuntimeError(f"Unresolved cross-references: {unresolved}")
    return source


def build_pandoc_source(original: str) -> tuple[str, list[TableSpec], str, int]:
    title = extract_command(original, "Title")
    authors = extract_command(original, "Author")
    address = extract_command(original, "address")
    correspondence = extract_command(original, "corres")
    simple_summary = extract_command(original, "simplesumm")
    abstract = extract_command(original, "abstract")
    keywords = extract_command(original, "keyword")

    body_start = original.index(r"\begin{document}") + len(r"\begin{document}")
    body_end = original.rindex(r"\end{document}")
    body = original[body_start:body_end]
    body = re.sub(r"\\bibliography\{[^}]+\}", "", body)
    body = re.sub(r"\\reftitle\{[^}]+\}", "", body)
    appendix_start = body.find(r"\begin{appendix}")
    if appendix_start >= 0:
        main_body = body[:appendix_start]
        appendix_body = body[appendix_start + len(r"\begin{appendix}") :]
        appendix_body = appendix_body.replace(r"\end{appendix}", "")
        appendix_index = 0

        def number_appendix(match: re.Match[str]) -> str:
            nonlocal appendix_index
            appendix_index += 1
            return f"\\section{{Appendix {chr(64 + appendix_index)}. {match.group(1)}}}"

        appendix_body = re.sub(r"\\section\{([^{}]+)\}", number_appendix, appendix_body)
        body = main_body + appendix_body
    body = body.replace(r"\appendixstart", "").replace(r"\end{appendix}", "")
    body = re.sub(r"\\vspace\{[^}]+\}", "", body)

    backmatter = (
        ("authorcontributions", "Author Contributions"),
        ("funding", "Funding"),
        ("dataavailability", "Data Availability Statement"),
        ("conflictsofinterest", "Conflicts of Interest"),
        ("acknowledgments", "Acknowledgments"),
    )
    for command, heading in backmatter:
        content = extract_command(body, command)
        start = body.index(f"\\{command}")
        brace = body.index("{", start)
        depth = 0
        end = None
        for index in range(brace, len(body)):
            if body[index] == "{":
                depth += 1
            elif body[index] == "}":
                depth -= 1
                if depth == 0:
                    end = index + 1
                    break
        if end is None:
            raise RuntimeError(f"Unclosed back-matter command: {command}")
        body = body[:start] + f"[[BACKMATTER|{heading}]] {content}\n\n" + body[end:]

    body, tables = extract_tables(body)
    body = replace_cross_references(body, original, tables)
    body = re.sub(r"\\subsubsection\{([^{}]+)\}", r"\\subsubsection{[[H3]] \1}", body)
    body = re.sub(r"\\subsection\{([^{}]+)\}", r"\\subsection{[[H2]] \1}", body)
    body = re.sub(r"\\section\{([^{}]+)\}", r"\\section{[[H1]] \1}", body)
    body = body.replace(r"$^{\circ}$", "°")
    body = body.replace(r"$^\circ$", "°")
    body = body.replace(r"\text{Butterfly index}", r"\mathrm{Butterfly\;index}")
    body = body.replace(r"\sqrt[3]{\Delta\mathrm{BI}}", "∛ΔBI")
    body = body.replace(r"\sqrt{\lvert \Delta\mathrm{BI}_{\max} \rvert}", "√|ΔBImax|")
    body = body.replace(r"\sqrt{|\Delta{BI}_{\max}|}", "√|ΔBImax|")

    addresses = [part.strip() for part in address.split(";") if part.strip()]
    address_block = "\n\n".join(f"[[ADDRESS]] {part}" for part in addresses)
    frontmatter = f"""
[[ARTICLE_TYPE]] Article

[[TITLE]] {title}

[[AUTHORS]] {authors}

{address_block}

[[CORRESPONDENCE]] {correspondence}

[[SIMPLE_SUMMARY]] Simple Summary

{simple_summary}

[[ABSTRACT]] Abstract

{abstract}

[[KEYWORDS]] Keywords: {keywords}

[[LINE]]
"""
    wrapper = "\n".join(
        (
            r"\documentclass{article}",
            r"\usepackage{amsmath}",
            r"\usepackage{graphicx}",
            r"\usepackage{float}",
            r"\usepackage{hyperref}",
            r"\begin{document}",
            frontmatter,
            body,
            r"\end{document}",
            "",
        )
    )
    expected_figures = len(re.findall(r"\\begin\{figure\}", original))
    return wrapper, tables, strip_latex(title), expected_figures


def set_repeat_table_header(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    marker = row_properties.find(qn("w:tblHeader"))
    if marker is None:
        marker = OxmlElement("w:tblHeader")
        row_properties.append(marker)
    marker.set(qn("w:val"), "1")


def keep_table_row_together(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    marker = row_properties.find(qn("w:cantSplit"))
    if marker is None:
        row_properties.append(OxmlElement("w:cantSplit"))


def set_cell_margins(cell, top: int = 70, start: int = 90, bottom: int = 70, end: int = 90) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    margins = cell_properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    properties = table._tbl.tblPr
    width_node = properties.first_child_found_in("w:tblW")
    if width_node is None:
        width_node = OxmlElement("w:tblW")
        properties.append(width_node)
    width_node.set(qn("w:w"), str(sum(widths)))
    width_node.set(qn("w:type"), "dxa")
    indent = properties.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), "0")
    indent.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        column_index = 0
        seen_cells: set[int] = set()
        for cell in row.cells:
            cell_key = id(cell._tc)
            if cell_key in seen_cells:
                continue
            seen_cells.add(cell_key)
            cell_properties = cell._tc.get_or_add_tcPr()
            span_node = cell_properties.first_child_found_in("w:gridSpan")
            span = int(span_node.get(qn("w:val"))) if span_node is not None else 1
            width = sum(widths[column_index : column_index + span])
            cell.width = width
            cell_width = cell_properties.first_child_found_in("w:tcW")
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                cell_properties.append(cell_width)
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            column_index += span


def scaled_widths(widths: list[int], total: int = 10466) -> list[int]:
    scale = total / sum(widths)
    result = [round(width * scale) for width in widths]
    result[-1] += total - sum(result)
    return result


def add_table_before(doc: Document, paragraph, spec: TableSpec, table_number: int) -> None:
    revised_labels = {
        "tab:deployment_metadata",
        "tab:harmonized_predictors",
        "tab:harmonized_templates",
        "tab:30min_candidates",
        "tab:nextday_candidates",
    }
    is_revised_table = spec.label in revised_labels
    display_number = (
        f"A{table_number}"
        if is_revised_table
        else (str(table_number) if table_number <= 7 else f"A{table_number - 7}")
    )
    caption = paragraph.insert_paragraph_before(f"Table {display_number}. {spec.caption}")
    caption.style = "MDPI_4.1_table_caption"
    caption.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=0, cols=len(spec.rows[0]))
    table.style = "MDPI_4.1_three_line_table"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row_index, values in enumerate(spec.rows):
        cells = table.add_row().cells
        keep_table_row_together(table.rows[-1])
        for column_index, value in enumerate(values):
            cell = cells[column_index]
            cell.text = value
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for cell_paragraph in cell.paragraphs:
                cell_paragraph.style = "MDPI_4.2_table_body"
                cell_paragraph.paragraph_format.space_before = Pt(0)
                cell_paragraph.paragraph_format.space_after = Pt(0)
                if is_revised_table:
                    cell_paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER if column_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    )
                elif table_number in {1, 5}:
                    cell_paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.LEFT if column_index == 1 else WD_ALIGN_PARAGRAPH.CENTER
                    )
                elif table_number in {2, 3, 6}:
                    cell_paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.LEFT if column_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    )
                elif table_number == 4:
                    cell_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif table_number == 7:
                    cell_paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.LEFT if column_index == 1 else WD_ALIGN_PARAGRAPH.CENTER
                    )
                else:
                    cell_paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER if column_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    )
                if row_index == 0:
                    for run in cell_paragraph.runs:
                        run.bold = True
        if row_index in spec.merge_rows and len(cells) > 1:
            merged = cells[0].merge(cells[-1])
            merged.text = values[0]
            for cell_paragraph in merged.paragraphs:
                cell_paragraph.style = "MDPI_4.2_table_body"
                cell_paragraph.paragraph_format.space_before = Pt(1.8)
                cell_paragraph.paragraph_format.space_after = Pt(1.8)
                for run in cell_paragraph.runs:
                    run.italic = True
    set_repeat_table_header(table.rows[0])
    columns = len(spec.rows[0])
    if spec.label == "tab:deployment_metadata":
        widths = scaled_widths([900, 2400, 1100, 1800, 1100])
    elif spec.label == "tab:harmonized_predictors":
        widths = scaled_widths([1200, 4100, 4100])
    elif spec.label == "tab:harmonized_templates":
        widths = scaled_widths([800, 2700, 6900])
    elif spec.label in {"tab:30min_candidates", "tab:nextday_candidates"}:
        widths = [1500, 8966]
    elif table_number >= 8 and columns == 2:
        widths = [1300, 9166]
    elif table_number == 4:
        widths = [3489, 3489, 3488]
    elif columns == 4:
        widths = scaled_widths([5010, 1420, 1420, 1420])
    elif table_number == 3:
        widths = scaled_widths([4900, 1000, 1000, 1200, 1170])
    elif columns == 5:
        widths = scaled_widths([1350, 4560, 1120, 1120, 1120])
    elif columns == 6:
        widths = scaled_widths([950, 4800, 880, 880, 880, 880])
    elif table_number in {1, 5}:
        widths = scaled_widths([850, 4120, 500, 1000, 900, 1000, 900])
    elif table_number == 7:
        widths = scaled_widths([700, 3300, 550, 950, 850, 950, 900, 1070])
    else:
        widths = [10466 // columns] * columns
        widths[-1] += 10466 - sum(widths)
    set_table_geometry(table, widths)
    paragraph._p.addprevious(table._tbl)
    paragraph._element.getparent().remove(paragraph._element)


def remove_marker(paragraph, marker: str) -> None:
    for run in paragraph.runs:
        if marker in run.text:
            run.text = run.text.replace(marker, "", 1)
            break
    for run in paragraph.runs:
        if run.text:
            run.text = run.text.lstrip()
            break


def bold_prefix(paragraph, prefix: str) -> None:
    for run in paragraph.runs:
        if not run.text:
            continue
        if run.text.startswith(prefix):
            tail = run.text[len(prefix) :]
            run.text = prefix
            run.bold = True
            if tail:
                new_run = paragraph.add_run(tail)
                new_run.bold = False
                run._r.addnext(new_run._r)
            return
    raise RuntimeError(f"Could not find prefix {prefix!r} in paragraph {paragraph.text!r}")


def metadata_table_from_template() -> OxmlElement:
    with zipfile.ZipFile(REFERENCE_TEMPLATE) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    tables = root.xpath(".//w:body/w:tbl", namespaces=namespace)
    if not tables:
        raise RuntimeError("The Insects template has no floating metadata table")
    table = deepcopy(tables[0])
    replacements = ["Academic Editor:", "Received:", "Revised:", "Accepted:", "Published:"]
    paragraphs = table.xpath(".//w:tr[1]/w:tc[1]/w:p", namespaces=namespace)
    if len(paragraphs) < len(replacements):
        raise RuntimeError("The Insects template metadata table is incomplete")
    for paragraph, replacement in zip(paragraphs, replacements):
        text_nodes = paragraph.xpath(".//w:t", namespaces=namespace)
        if not text_nodes:
            raise RuntimeError("A metadata paragraph has no text node")
        text_nodes[0].text = replacement
        for text_node in text_nodes[1:]:
            text_node.text = ""
    return table


def add_update_fields_setting(doc: Document) -> None:
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def style_document(
    doc: Document,
    tables: list[TableSpec],
    source_commit: str,
    manuscript_title: str,
    expected_figures: int,
) -> None:
    author_paragraph = next(
        paragraph for paragraph in doc.paragraphs if paragraph.text.startswith("[[AUTHORS]]")
    )
    author_paragraph._p.addnext(metadata_table_from_template())
    add_update_fields_setting(doc)

    frontmatter_styles = {
        "[[ARTICLE_TYPE]]": "MDPI_1.1_article_type",
        "[[TITLE]]": "MDPI_1.2_title",
        "[[AUTHORS]]": "MDPI_1.3_authornames",
        "[[ADDRESS]]": "MDPI_1.6_affiliation",
        "[[CORRESPONDENCE]]": "MDPI_1.6_affiliation",
        "[[SIMPLE_SUMMARY]]": "MDPI_1.7_abstract",
        "[[ABSTRACT]]": "MDPI_1.7_abstract",
        "[[KEYWORDS]]": "MDPI_1.8_keywords",
        "[[LINE]]": "MDPI_1.9_line",
    }
    body_started = False
    references_started = False
    current_main_section = 0
    current_appendix = None
    subsection = 0
    subsubsection = 0
    previous_was_heading = False
    expect_figure_caption = False
    expect_abstract_text = False
    figure_number = 1

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        matched_frontmatter = False
        for marker, style_name in frontmatter_styles.items():
            if text.startswith(marker):
                remove_marker(paragraph, marker)
                paragraph.style = style_name
                matched_frontmatter = True
                if marker in {"[[SIMPLE_SUMMARY]]", "[[ABSTRACT]]"}:
                    for run in paragraph.runs:
                        run.bold = True
                    expect_abstract_text = True
                elif marker == "[[KEYWORDS]]":
                    bold_prefix(paragraph, "Keywords:")
                elif marker == "[[CORRESPONDENCE]]" and paragraph.text.startswith(
                    "Correspondence:"
                ):
                    paragraph.text = f"*\t{paragraph.text}"
                    bold_prefix(paragraph, "*")
                break
        if matched_frontmatter:
            continue

        text = paragraph.text.strip()
        if expect_abstract_text and text:
            paragraph.style = "MDPI_1.7_abstract"
            expect_abstract_text = False
            continue
        if text.startswith("[[H1]]"):
            body_started = True
            remove_marker(paragraph, "[[H1]]")
            heading = paragraph.text.strip()
            if heading.startswith("Appendix "):
                current_appendix = heading.split()[1].rstrip(".")
                subsection = 0
                subsubsection = 0
            else:
                current_appendix = None
                current_main_section += 1
                subsection = 0
                subsubsection = 0
                paragraph.text = f"{current_main_section}. {heading}"
            paragraph.style = "MDPI_2.1_heading1"
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
            previous_was_heading = True
            continue
        if text.startswith("[[H2]]"):
            remove_marker(paragraph, "[[H2]]")
            subsection += 1
            subsubsection = 0
            section_prefix = current_appendix or current_main_section
            paragraph.text = f"{section_prefix}.{subsection}. {paragraph.text.strip()}"
            paragraph.style = "MDPI_2.2_heading2"
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
            previous_was_heading = True
            continue
        if text.startswith("[[H3]]"):
            remove_marker(paragraph, "[[H3]]")
            subsubsection += 1
            paragraph.text = (
                f"{current_appendix or current_main_section}.{subsection}.{subsubsection}. {paragraph.text.strip()}"
            )
            paragraph.style = "MDPI_2.3_heading3"
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
            previous_was_heading = True
            continue
        if text.startswith("[[BACKMATTER|"):
            marker_match = re.match(r"\[\[BACKMATTER\|([^]]+)\]\]", text)
            if not marker_match:
                raise RuntimeError(f"Malformed back-matter marker: {text}")
            label = marker_match.group(1)
            remove_marker(paragraph, marker_match.group(0))
            paragraph.text = f"{label}: {paragraph.text.strip()}"
            paragraph.style = "MDPI_6.2_back_matter"
            bold_prefix(paragraph, f"{label}:")
            previous_was_heading = False
            continue
        if text == "References":
            paragraph.style = "MDPI_2.1_heading1"
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
            references_started = True
            previous_was_heading = True
            continue
        if references_started:
            paragraph.style = "MDPI_8.1_references"
            for text_node in paragraph._p.xpath(".//w:t"):
                if text_node.text and text_node.text.strip():
                    text_node.text = re.sub(r"^\s*\d+\.\s+", "", text_node.text, count=1)
                    break
            previous_was_heading = False
            continue

        if expect_figure_caption and text:
            run = OxmlElement("w:r")
            text_node = OxmlElement("w:t")
            text_node.set(qn("xml:space"), "preserve")
            text_node.text = f"Figure {figure_number}. "
            run.append(text_node)
            children = list(paragraph._p)
            insert_index = 1 if children and children[0].tag == qn("w:pPr") else 0
            paragraph._p.insert(insert_index, run)
            paragraph.style = "MDPI_5.1_figure_caption"
            figure_number += 1
            expect_figure_caption = False
            previous_was_heading = False
            continue
        if "w:drawing" in paragraph._p.xml:
            paragraph.style = "MDPI_5.2_figure"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            expect_figure_caption = True
            previous_was_heading = False
            continue
        if body_started and not text.startswith("[[TABLE_"):
            paragraph.style = (
                "MDPI_3.2_text_no_indent" if previous_was_heading else "MDPI_3.1_text"
            )
            previous_was_heading = False

    table_by_marker = {table.marker: table for table in tables}
    for paragraph in list(doc.paragraphs):
        marker = paragraph.text.strip()
        if marker in table_by_marker:
            spec = table_by_marker[marker]
            add_table_before(doc, paragraph, spec, tables.index(spec) + 1)

    if figure_number - 1 != expected_figures:
        raise RuntimeError(
            f"Expected {expected_figures} figures, found {figure_number - 1}"
        )
    for shape in doc.inline_shapes:
        max_width = Inches(6.7)
        max_height = Inches(7.5)
        scale = min(1.0, max_width / shape.width, max_height / shape.height)
        shape.width = int(shape.width * scale)
        shape.height = int(shape.height * scale)

    doc.core_properties.title = manuscript_title
    doc.core_properties.subject = f"MDPI Insects manuscript submitted from Git commit {source_commit}"
    doc.core_properties.author = "Kyle Nessen, Peter C. Ibsen, Jay E. Diffendorfer, Francis X. Villablanca"
    doc.core_properties.last_modified_by = ""
    doc.core_properties.keywords = "monarch butterfly; wind disruption; thermoregulation; MDPI Insects"


def strip_review_parts(path: Path) -> None:
    word_namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    relationship_namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    content_type_namespace = "http://schemas.openxmlformats.org/package/2006/content-types"
    comment_relationship_types = {
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments",
        "http://schemas.microsoft.com/office/2011/relationships/commentsExtended",
    }
    custom_properties_relationship_type = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties"
    )
    hyperlink_relationship_type = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
    )
    removable = {
        "word/comments.xml",
        "word/commentsExtended.xml",
        "word/_rels/footnotes.xml.rels",
        "docProps/custom.xml",
    }
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False, dir=path.parent) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with zipfile.ZipFile(path) as source:
            overrides: dict[str, bytes] = {}
            for story in ["word/document.xml"] + [
                name
                for name in source.namelist()
                if re.fullmatch(r"word/(?:header|footer)\d+\.xml", name)
            ]:
                root = etree.fromstring(source.read(story))
                changed = False
                for tag in ("commentRangeStart", "commentRangeEnd", "commentReference"):
                    for element in root.xpath(f".//w:{tag}", namespaces={"w": word_namespace}):
                        element.getparent().remove(element)
                        changed = True
                if changed:
                    overrides[story] = etree.tostring(
                        root, xml_declaration=True, encoding="UTF-8", standalone="yes"
                    )

            rels_path = "word/_rels/document.xml.rels"
            rels_root = etree.fromstring(source.read(rels_path))
            rels_changed = False
            for relationship in list(rels_root.findall(f"{{{relationship_namespace}}}Relationship")):
                if relationship.get("Type") in comment_relationship_types:
                    rels_root.remove(relationship)
                    rels_changed = True
            relationship_ids = {
                relationship.get("Id")
                for relationship in rels_root.findall(
                    f"{{{relationship_namespace}}}Relationship"
                )
            }
            document_root = etree.fromstring(source.read("word/document.xml"))
            creative_commons_links = document_root.xpath(
                ".//w:hyperlink[@r:id='rId8']",
                namespaces={
                    "w": word_namespace,
                    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
                },
            )
            if creative_commons_links and "rId8" not in relationship_ids:
                relationship = etree.Element(
                    f"{{{relationship_namespace}}}Relationship",
                    Id="rId8",
                    Type=hyperlink_relationship_type,
                    Target="https://creativecommons.org/licenses/by/4.0/",
                    TargetMode="External",
                )
                rels_root.append(relationship)
                rels_changed = True
            if rels_changed:
                overrides[rels_path] = etree.tostring(
                    rels_root, xml_declaration=True, encoding="UTF-8", standalone="yes"
                )

            package_rels_path = "_rels/.rels"
            package_rels_root = etree.fromstring(source.read(package_rels_path))
            package_rels_changed = False
            for relationship in list(
                package_rels_root.findall(f"{{{relationship_namespace}}}Relationship")
            ):
                if relationship.get("Type") == custom_properties_relationship_type:
                    package_rels_root.remove(relationship)
                    package_rels_changed = True
            if package_rels_changed:
                overrides[package_rels_path] = etree.tostring(
                    package_rels_root, xml_declaration=True, encoding="UTF-8", standalone="yes"
                )

            content_types_path = "[Content_Types].xml"
            content_types_root = etree.fromstring(source.read(content_types_path))
            content_types_changed = False
            for override in list(content_types_root.findall(f"{{{content_type_namespace}}}Override")):
                if override.get("PartName") in (
                    "/word/comments.xml",
                    "/word/commentsExtended.xml",
                    "/docProps/custom.xml",
                ):
                    content_types_root.remove(override)
                    content_types_changed = True
            if content_types_changed:
                overrides[content_types_path] = etree.tostring(
                    content_types_root, xml_declaration=True, encoding="UTF-8", standalone="yes"
                )

            with zipfile.ZipFile(temporary_path, "w", zipfile.ZIP_DEFLATED) as target:
                for item in source.infolist():
                    if item.filename in removable:
                        continue
                    target.writestr(item, overrides.get(item.filename, source.read(item.filename)))
        temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)


def make_reference_docx(path: Path) -> None:
    with zipfile.ZipFile(REFERENCE_TEMPLATE) as source, zipfile.ZipFile(
        path, "w", zipfile.ZIP_DEFLATED
    ) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    TEMPLATE_MAIN_CONTENT_TYPE.encode(), DOCUMENT_MAIN_CONTENT_TYPE.encode()
                )
            target.writestr(item, data)


def validate(
    path: Path,
    tables: list[TableSpec],
    source_commit: str,
    expected_figures: int,
    strict_submitted: bool,
) -> None:
    doc = Document(path)
    all_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    headings = [
        paragraph.text
        for paragraph in doc.paragraphs
        if paragraph.style and paragraph.style.name == "MDPI_2.1_heading1"
    ]
    required = {
        "1. Introduction",
        "2. Materials and Methods",
        "3. Results",
        "4. Discussion",
        "5. Conclusions",
        "References",
    }
    if strict_submitted:
        required.update(
            {
                "Appendix A. 30-Minute Wind Disruption Analysis: Candidate Models",
                "Appendix B. Threshold Wind Disruption Analysis: Candidate Models",
                "Appendix C. Site Fidelity Analysis (Next Day Window): Candidate Models",
                "Appendix D. 24-Hour Robustness Analysis: Candidate Models",
            }
        )
    else:
        required.update(
            {
                "Appendix A. Monitoring Deployment Metadata",
                "Appendix B. Candidate Model Strategy",
            }
        )
    missing = required.difference(headings)
    if missing:
        raise RuntimeError(f"Missing top-level sections: {sorted(missing)}")
    if len(doc.inline_shapes) != expected_figures:
        raise RuntimeError(
            f"Expected {expected_figures} figures, found {len(doc.inline_shapes)}"
        )
    if len(doc.tables) != len(tables) + 1:
        raise RuntimeError(
            f"Expected {len(tables) + 1} tables including metadata, found {len(doc.tables)}"
        )
    bibliography = [
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.style and paragraph.style.name == "MDPI_8.1_references"
    ]
    if strict_submitted and len(bibliography) != 66:
        raise RuntimeError(f"Expected 66 bibliography entries, found {len(bibliography)}")
    if not bibliography:
        raise RuntimeError("No bibliography entries were exported")
    for label in (
        "Author Contributions:",
        "Funding:",
        "Data Availability Statement:",
        "Acknowledgments:",
        "Conflicts of Interest:",
    ):
        if label not in all_text:
            raise RuntimeError(f"Missing back-matter label: {label}")
    if re.search(r"\\(?:cite|ref)\{|\[\[(?:ARTICLE|TITLE|TABLE|AUTHORS|H[123]|BACKMATTER)", all_text):
        raise RuntimeError("Raw source markers remain in the exported document")
    forbidden_text = (
        "How to Use This Template",
        "Type of the Paper (Article, Review, Communication, etc.)",
        "Firstname Lastname",
        "keyword 1; keyword 2; keyword 3",
    )
    for value in forbidden_text:
        if value in all_text:
            raise RuntimeError(f"Template instructional text remains: {value}")
    if f"Git commit {source_commit}" not in doc.core_properties.subject:
        raise RuntimeError("Source commit metadata is missing")
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        required_parts = {
            "word/header1.xml",
            "word/header2.xml",
            "word/header3.xml",
            "word/footer1.xml",
            "word/footer2.xml",
            "word/media/image3.png",
            "word/media/image4.png",
            "word/styles.xml",
            "word/settings.xml",
        }
        missing_parts = required_parts.difference(names)
        if missing_parts:
            raise RuntimeError(f"Template package parts are missing: {sorted(missing_parts)}")
        content_types = archive.read("[Content_Types].xml")
        if DOCUMENT_MAIN_CONTENT_TYPE.encode() not in content_types:
            raise RuntimeError("The output package is not a standard DOCX")
        if TEMPLATE_MAIN_CONTENT_TYPE.encode() in content_types:
            raise RuntimeError("The output package still declares a Word template")
        forbidden_parts = {name for name in names if "comment" in name.lower() or name.endswith("people.xml")}
        if forbidden_parts:
            raise RuntimeError(f"Review parts remain: {sorted(forbidden_parts)}")
        if "word/_rels/footnotes.xml.rels" in names:
            raise RuntimeError("Unused footnote relationships remain")
        if "docProps/custom.xml" in names:
            raise RuntimeError("Conversion-tool custom properties remain")
        package_relationships = archive.read("_rels/.rels")
        if b"custom-properties" in package_relationships:
            raise RuntimeError("Custom-properties package relationship remains")
        document_xml = archive.read("word/document.xml")
        for marker in (b"<w:ins", b"<w:del", b"<w:moveFrom", b"<w:moveTo", b"commentRange"):
            if marker in document_xml:
                raise RuntimeError(f"Review markup remains: {marker.decode(errors='ignore')}")
        document_root = etree.fromstring(document_xml)
        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        relationship_namespace = (
            "http://schemas.openxmlformats.org/package/2006/relationships"
        )
        office_relationship_namespace = (
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        )
        relationship_root = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
        declared_relationships = {
            relationship.get("Id")
            for relationship in relationship_root.findall(
                f"{{{relationship_namespace}}}Relationship"
            )
        }
        referenced_relationships = set(
            document_root.xpath(
                "//@r:id | //@r:embed | //@r:link",
                namespaces={"r": office_relationship_namespace},
            )
        )
        missing_relationships = referenced_relationships.difference(declared_relationships)
        if missing_relationships:
            raise RuntimeError(
                f"Document relationships are missing: {sorted(missing_relationships)}"
            )
        sections = document_root.xpath(".//w:sectPr", namespaces=namespace)
        if len(sections) != 1:
            raise RuntimeError(f"Expected one section, found {len(sections)}")
        section = sections[0]
        page_size = section.find(qn("w:pgSz"))
        margins = section.find(qn("w:pgMar"))
        line_numbers = section.find(qn("w:lnNumType"))
        if page_size is None or (page_size.get(qn("w:w")), page_size.get(qn("w:h"))) != (
            "11906",
            "16838",
        ):
            raise RuntimeError("The A4 template page size was not preserved")
        expected_margins = {
            "top": "1417",
            "right": "720",
            "bottom": "907",
            "left": "720",
            "header": "720",
            "footer": "612",
        }
        if margins is None or any(
            margins.get(qn(f"w:{name}")) != value for name, value in expected_margins.items()
        ):
            raise RuntimeError("The template margins were not preserved")
        if line_numbers is None or line_numbers.get(qn("w:countBy")) != "1":
            raise RuntimeError("Continuous line numbering was not preserved")
        if section.find(qn("w:titlePg")) is None:
            raise RuntimeError("The distinct first-page layout was not preserved")
        header_text = archive.read("word/header2.xml")
        if b"PAGE" not in header_text or b"NUMPAGES" not in header_text:
            raise RuntimeError("The running header page fields are missing")


def export(commit: str, output: Path, pandoc: str) -> None:
    resolved = run("git", "rev-parse", commit, cwd=REPO_ROOT, capture=True)
    strict_submitted = resolved == EXPECTED_COMMIT
    if sha256(REFERENCE_TEMPLATE) != EXPECTED_TEMPLATE_SHA256:
        raise RuntimeError("The MDPI Insects template does not match its recorded checksum")

    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="monarch-submitted-docx-") as temp_name:
        temp_dir = Path(temp_name)
        source_dir = temp_dir / "source"
        source_dir.mkdir()
        archive_path = temp_dir / "source.tar"
        with archive_path.open("wb") as archive_handle:
            subprocess.run(("git", "archive", resolved), cwd=REPO_ROOT, check=True, stdout=archive_handle)
        run("tar", "-xf", str(archive_path), "-C", str(source_dir))

        original = (source_dir / "manuscript.tex").read_text(encoding="utf-8")
        pandoc_source, tables, manuscript_title, expected_figures = build_pandoc_source(original)
        wrapped_tex = temp_dir / "manuscript-word.tex"
        raw_docx = temp_dir / "manuscript-word-raw.docx"
        reference_docx = temp_dir / "insects-reference.docx"
        wrapped_tex.write_text(pandoc_source, encoding="utf-8")
        make_reference_docx(reference_docx)

        run(
            pandoc,
            str(wrapped_tex),
            "--from=latex",
            "--to=docx",
            "--citeproc",
            f"--csl={CSL}",
            f"--bibliography={source_dir / 'bibliography' / 'references.bib'}",
            f"--resource-path={source_dir}",
            f"--reference-doc={reference_docx}",
            "--metadata=link-citations:false",
            "--metadata=reference-section-title:References",
            f"--output={raw_docx}",
            cwd=source_dir,
        )

        doc = Document(raw_docx)
        style_document(doc, tables, resolved, manuscript_title, expected_figures)
        doc.save(output)
        strip_review_parts(output)
    validate(output, tables, resolved, expected_figures, strict_submitted)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--commit", default="main")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--pandoc", default=shutil.which("pandoc") or "pandoc")
    return parser.parse_args()


if __name__ == "__main__":
    arguments = parse_args()
    export(arguments.commit, arguments.output.resolve(), arguments.pandoc)
