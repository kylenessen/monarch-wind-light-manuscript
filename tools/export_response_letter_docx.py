#!/usr/bin/env python3
"""Export the finalized Typst reviewer response as a clean Word document."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TITLE = "Point-by-Point Response"
SUBTITLE = "Response to the Academic Editor and Reviewers"
MANUSCRIPT_TITLE = (
    "Wind Associations with Overwintering Monarch Butterfly Cluster Size "
    "Depend on Temperature and Sun Exposure"
)


def set_run_font(run, name: str, size: float, color: str = "000000", bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_paragraph_shading(paragraph, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_left_border(paragraph, color: str = "2F6F70", size: str = "18") -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    left = borders.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        borders.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), "7")
    left.set(qn("w:color"), color)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend((begin, instruction, separate, text, end))
    set_run_font(run, "Calibri", 9, "666666")


def parse_bracket_group(source: str, start: int) -> tuple[str, int]:
    if source[start] != "[":
        raise ValueError(f"Expected '[' at offset {start}")
    depth = 0
    quoted = False
    escaped = False
    for index in range(start, len(source)):
        char = source[index]
        if quoted:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == '"':
                quoted = False
            continue
        if char == '"':
            quoted = True
        elif char == "[":
            depth += 1
        elif char == "]":
            depth -= 1
            if depth == 0:
                return source[start + 1 : index], index + 1
    raise ValueError("Unclosed bracket group")


def clean_typst(value: str) -> str:
    value = re.sub(
        r'#link\("([^"]+)"\)\[([^]]+)\]',
        lambda match: f"{match.group(2)} ({match.group(1)})",
        value,
    )
    value = value.replace("\\\n", " ")
    return re.sub(r"[ \t]+", " ", value).strip()


def parse_entries(source: str) -> list[tuple[str, str, str]]:
    entries: list[tuple[str, str, str]] = []
    position = 0
    while True:
        marker = source.find("#entry(", position)
        if marker < 0:
            break
        cursor = marker + len("#entry(")
        groups: list[str] = []
        for _ in range(3):
            while cursor < len(source) and source[cursor] in " \t\r\n,":
                cursor += 1
            group, cursor = parse_bracket_group(source, cursor)
            groups.append(clean_typst(group))
        entries.append((groups[0], groups[1], groups[2]))
        position = cursor
    if len(entries) != 23:
        raise RuntimeError(f"Expected 23 response entries, found {len(entries)}")
    return entries


def add_body_paragraph(doc: Document, text: str, *, bold_prefix: str | None = None):
    paragraph = doc.add_paragraph()
    paragraph.style = "Normal"
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    if bold_prefix and text.startswith(bold_prefix):
        lead = paragraph.add_run(bold_prefix)
        set_run_font(lead, "Calibri", 11, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix) :])
        set_run_font(rest, "Calibri", 11)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, "Calibri", 11)
    return paragraph


def add_entry(doc: Document, title: str, concern: str, answer: str) -> None:
    heading = doc.add_paragraph(style="Heading 2")
    heading.paragraph_format.keep_with_next = True
    heading.add_run(title)

    comment = doc.add_paragraph()
    comment.paragraph_format.left_indent = Inches(0.05)
    comment.paragraph_format.right_indent = Inches(0.05)
    comment.paragraph_format.space_before = Pt(0)
    comment.paragraph_format.space_after = Pt(4)
    comment.paragraph_format.line_spacing = 1.1
    comment.paragraph_format.keep_with_next = True
    set_paragraph_shading(comment, "F2F5F7")
    lead = comment.add_run("Comment. ")
    set_run_font(lead, "Calibri", 11, bold=True)
    body = comment.add_run(concern)
    set_run_font(body, "Calibri", 11)

    answer_parts = [part.strip() for part in answer.split("\n\n") if part.strip()]
    for index, part in enumerate(answer_parts):
        response = doc.add_paragraph()
        response.paragraph_format.left_indent = Inches(0.12)
        response.paragraph_format.space_before = Pt(0)
        response.paragraph_format.space_after = Pt(6 if index == len(answer_parts) - 1 else 4)
        response.paragraph_format.line_spacing = 1.1
        set_left_border(response)
        if index == 0:
            lead = response.add_run("Response. ")
            set_run_font(lead, "Calibri", 11, "1D5051", bold=True)
        body = response.add_run(part)
        set_run_font(body, "Calibri", 11)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def export(source: Path, output: Path) -> None:
    typst = source.read_text(encoding="utf-8")
    entries = parse_entries(typst)
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)
    add_page_field(section.footer.paragraphs[0])

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run(TITLE)
    set_run_font(run, "Calibri", 20, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run(SUBTITLE)
    set_run_font(run, "Calibri", 13, "444444")

    metadata = doc.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_after = Pt(14)
    lead = metadata.add_run("Revised manuscript title\n")
    set_run_font(lead, "Calibri", 10.5, bold=True)
    manuscript = metadata.add_run(MANUSCRIPT_TITLE + "\n")
    set_run_font(manuscript, "Calibri", 10.5, italic=True)
    journal = metadata.add_run("Journal: Insects")
    set_run_font(journal, "Calibri", 10.5, bold=True)

    add_body_paragraph(doc, "Dear Academic Editor and Reviewers,")
    add_body_paragraph(
        doc,
        "Thank you for the careful and constructive reviews. The manuscript is substantially "
        "stronger because of this feedback. We shortened the paper and revised its analyses, "
        "interpretation, and scope. The revised title and "
        "all summary sections now describe conditional associations under the monitored "
        "conditions rather than a general or causal rejection of wind effects. We removed "
        "unsupported analyses and management recommendations, clarified the environmental and "
        "image-based measurements, repeated the candidate-model comparisons, added sensitivity "
        "analyses, and made the processed data and analytical materials public. The responses "
        "below address every comment. Manuscript sections are named directly because line "
        "numbers may change during journal production.",
    )

    academic_editor_count = 3
    reviewer_one_count = 16
    heading = doc.add_paragraph(style="Heading 1")
    heading.add_run("Academic Editor")
    for index, entry in enumerate(entries):
        if index == academic_editor_count:
            heading = doc.add_paragraph(style="Heading 1")
            heading.add_run("Reviewer 1")
        elif index == academic_editor_count + reviewer_one_count:
            heading = doc.add_paragraph(style="Heading 1")
            heading.add_run("Reviewer 2")
        add_entry(doc, *entry)

    add_body_paragraph(
        doc,
        "We again thank the Academic Editor and both reviewers. Their comments led to a shorter, "
        "more transparent manuscript with narrower claims, corrected model comparisons, clearer "
        "measurement limitations, and stronger reproducibility materials.",
    )
    add_body_paragraph(doc, "Sincerely,")
    add_body_paragraph(
        doc,
        "Kyle Nessen, Peter C. Ibsen, Jay E. Diffendorfer, and Francis X. Villablanca",
    )

    doc.core_properties.title = "Response to the Academic Editor and Reviewers"
    doc.core_properties.subject = MANUSCRIPT_TITLE
    doc.core_properties.author = (
        "Kyle Nessen, Peter C. Ibsen, Jay E. Diffendorfer, Francis X. Villablanca"
    )
    doc.core_properties.last_modified_by = ""
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=Path("reviewers/response-letter.typ"))
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    export(args.source.resolve(), args.output.resolve())
